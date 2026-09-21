#!/usr/bin/env python3
"""
WirForce BYOC 座位表解析器
把「地圖式」Google Sheets 座位表（下載成 .xlsx）轉成 seats.json。

版面規則（依座位表維護者的慣例）：
  - 座位編號格：內容符合 A01 / B123 / C466 這種格式
  - 同一列（試算表的一個 row）會同時橫跨 A/B/C 三區，各區的排法互相獨立，
    所以先把每一列切成「區塊」（欄位相鄰、同一區的一串座位）再逐塊判斷方向
  - 每個區塊旁邊的走道欄位寫著「面上」或「面下」：
      面上 = 這排人坐在桌子上緣，暱稱寫在編號的下一列
      面下 = 這排人坐在桌子下緣，暱稱寫在編號的上一列
  - 暱稱儲存格若合併跨越兩三個座位 = 大桌（一人多位）
  - 兩個區塊上下緊貼且面向相反 = 背對背；中間夾著暱稱列 = 同一張桌子面對面

用法：
  python build_seats.py 座位表.xlsx               # 讀本機檔案
  python build_seats.py --sheet-id <ID>          # 直接從 Google 下載（試算表需公開檢視）
輸出：seats.json（與 index.html 放在同一個資料夾）
"""
import argparse
import json
import re
import sys
import tempfile
import urllib.request
from datetime import datetime, timezone, timedelta

import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

SEAT_RE = re.compile(r"^\s*([A-Z])\s*0*(\d{1,4})\s*$")
# 「空白領域N 團名」：沒有座位編號的營區，單獨列成 zones，不混進暱稱
ZONE_RE = re.compile(r"^空白領域\s*0*(\d{1,2})\s*[:：\-]?\s*(.*)$")
FACING = {"面上": +1, "面下": -1}  # 值 = 暱稱寫在編號的哪一側（+1 下方、-1 上方）
MIN_SEATS_PER_ID_ROW = 3  # 一列至少要有幾個座位編號才算「座位編號列」
MAX_NAME_GAP = 2          # 暱稱最多離編號列幾列（正常 1，原表偶爾多空一列）
MAX_TABLE_GAP = 10        # 同桌對面最多隔幾列（原表會在中間插入「社交走道」說明，隔很開）
MAX_BIG_SEATS = 3         # 一格暱稱最多涵蓋幾個座位（再寬就是走道說明，不是大桌）
TZ_TAIPEI = timezone(timedelta(hours=8))


def download_xlsx(sheet_id: str) -> str:
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    req = urllib.request.Request(url, headers={"User-Agent": "wirforce-seat-finder"})
    with urllib.request.urlopen(req, timeout=60) as r:
        tmp.write(r.read())
    tmp.close()
    return tmp.name


# 「在原座位表中查看」要連到正確的分頁。gid 不在 xlsx 裡（那是 Google 自己的概念），
# 但 htmlview 端點的 bootstrap 資料有 name -> gid 的對照。
GID_RE = re.compile(r'items\.push\(\{name:\s*"([^"]+)"[^}]*?gid:\s*"(\d+)"')


def fetch_gids(sheet_id: str) -> dict:
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/htmlview"
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=30) as r:
            html = r.read().decode("utf-8", "replace")
        return {name: gid for name, gid in GID_RE.findall(html)}
    except Exception:
        return {}      # 抓不到就算了，連結會退成只開試算表不跳格子


def cell_text(v) -> str:
    if v is None:
        return ""
    return str(v).replace("\n", " ").strip()


def fill_hex(cell):
    """讀座位格底色（讓網頁的小地圖顏色跟原表一致），讀不到就回 None。"""
    try:
        c = cell.fill.fgColor
        if cell.fill.fill_type and c is not None and c.type == "rgb" and c.rgb:
            rgb = str(c.rgb)[-6:]
            if rgb.upper() not in ("000000", "FFFFFF"):
                return "#" + rgb
    except Exception:
        pass
    return None


class Grid:
    """把工作表整理成 (row, col) -> 文字，並記錄每個合併儲存格的範圍。"""

    def __init__(self, ws):
        self.ws = ws
        self.text = {}
        self.span = {}  # (row, col) of anchor -> (min_col, max_col, n_rows)
        for rng in ws.merged_cells.ranges:
            self.span[(rng.min_row, rng.min_col)] = (
                rng.min_col, rng.max_col, rng.max_row - rng.min_row + 1)
        for row in ws.iter_rows():
            for c in row:
                t = cell_text(c.value)
                if t:
                    self.text[(c.row, c.column)] = t

    def col_span(self, r, c):
        s = self.span.get((r, c))
        return (s[0], s[1]) if s else (c, c)

    def row_span(self, r, c):
        s = self.span.get((r, c))
        return s[2] if s else 1


def find_seats(grid: Grid):
    seats = []
    for (r, c), t in grid.text.items():
        m = SEAT_RE.match(t)
        if m:
            area, num = m.group(1), int(m.group(2))
            c0, c1 = grid.col_span(r, c)
            seats.append({
                "id": f"{area}{num:04d}",   # 一律 1 英文 + 4 數字，不足補 0
                "area": area, "num": num, "r": r, "c0": c0, "c1": c1,
                "cell": f"{get_column_letter(c)}{r}",
                "color": fill_hex(grid.ws.cell(row=r, column=c)),
            })
    return seats


def find_zones(grid: Grid):
    """空白領域：14 個營區，沒有座位編號，只有編號與團名。"""
    zones = {}
    for (r, c), t in grid.text.items():
        m = ZONE_RE.match(t)
        if not m:
            continue
        no, name = int(m.group(1)), m.group(2).strip()
        c0, c1 = grid.col_span(r, c)
        z = {"no": no, "name": name, "label": f"空白領域{no:02d}",
             "cell": f"{get_column_letter(c)}{r}", "color": fill_hex(grid.ws.cell(row=r, column=c)),
             "r": r, "r1": r + grid.row_span(r, c) - 1, "c0": c0, "c1": c1}
        if no not in zones or (name and not zones[no]["name"]):
            zones[no] = z
    return [zones[k] for k in sorted(zones)]


def zone_notes(grid: Grid, zones):
    """營區介紹：寫在「空白領域N 團名」那格的正下方。原表多半還沒寫，有多少算多少。"""
    for z in zones:
        col = column_index_from_string("".join(ch for ch in z["cell"] if ch.isalpha()))
        t = grid.text.get((z["r1"] + 1, col)) or grid.text.get((z["r"] + 1, col))
        if t and not ZONE_RE.match(t) and not SEAT_RE.match(t):
            z["about"] = t
    return zones


def find_notes(wb):
    """座位備註：某張工作表有「編號／ID／備註」三欄，玩家在自己座位旁邊留話。"""
    for ws in wb.worksheets:
        head = None
        for r in range(1, min(6, ws.max_row) + 1):
            row = {c: cell_text(ws.cell(row=r, column=c).value) for c in range(1, ws.max_column + 1)}
            if "編號" in row.values() and "備註" in row.values():
                head = (r, row)
                break
        if not head:
            continue
        r0, row = head
        groups = [c for c, t in row.items() if t == "編號"]
        notes = []
        for r in range(r0 + 1, ws.max_row + 1):
            for c in groups:
                note = cell_text(ws.cell(row=r, column=c + 2).value)
                if not note:
                    continue
                m = SEAT_RE.match(cell_text(ws.cell(row=r, column=c).value))
                notes.append({
                    "seat": f"{m.group(1)}{int(m.group(2)):04d}" if m else None,
                    "name": cell_text(ws.cell(row=r, column=c + 1).value),
                    "text": note,
                    "cell": f"{get_column_letter(c + 2)}{r}",
                })
        if notes:
            return notes, ws.title
    return [], None


# 場地裡的吃喝與設施。原表沒有欄位標記這些，只能靠關鍵字認，所以會把抓到的印出來人工核對。
FOOD_RE = re.compile(r"食堂|咖啡|cafe|星巴克|7-?11|超商|便利商店|餐廳|飲料|whiskey|ktv|居酒屋", re.I)
FACILITY_RE = re.compile(r"^wc$|廁所|洗手間|郵局|atm|出入口|吸菸|置物", re.I)
HOURS_RE = re.compile(r"\d{1,2}[:：]\d{2}\s*[–\-~〜]\s*\d{1,2}[:：]\d{2}")
# 維護者的閒聊也常提到星巴克、7-11，但那是句子不是地點標籤 —— 有句讀就不是招牌
CHATTY_RE = re.compile(r"[。，？！（）()]")


def find_places(grid: Grid, seat_cols):
    """挑出吃喝與設施的標籤。座位帶裡面的不算（那是暱稱），只看場地圖其他地方。"""
    lo, hi = seat_cols
    found = {}
    for (r, c), t in sorted(grid.text.items()):
        if lo <= c <= hi or len(t) > 100:
            continue
        if CHATTY_RE.search(t):
            continue
        kind = "food" if FOOD_RE.search(t) else ("facility" if FACILITY_RE.search(t) else None)
        if not kind:
            continue
        t = t.strip('"「」 ')
        hours = HOURS_RE.findall(t)
        key = re.sub(r"\s+", "", t)
        # 同一個地點在「攤位進駐表」和場地圖上各出現一次，留資訊比較多的那個
        if key in found and len(found[key]["hours"]) >= len(hours):
            continue
        found[key] = {"name": t, "kind": kind, "hours": hours,
                      "cell": f"{get_column_letter(c)}{r}",
                      "color": fill_hex(grid.ws.cell(row=r, column=c))}
    return sorted(found.values(), key=lambda p: (p["kind"] != "food", p["name"]))


def split_blocks(row_seats):
    """把同一列的座位切成區塊：欄位不相鄰或換區就切開。"""
    blocks, cur = [], []
    for s in sorted(row_seats, key=lambda s: s["c0"]):
        if cur and (s["c0"] - cur[-1]["c1"] > 1 or s["area"] != cur[-1]["area"]):
            blocks.append(cur)
            cur = []
        cur.append(s)
    if cur:
        blocks.append(cur)
    return blocks


def block_distance(block, col):
    if block["c0"] <= col <= block["c1"]:
        return 0
    return block["c0"] - col if col < block["c0"] else col - block["c1"]


def mark_reader(grid: Grid, blocks, blocks_by_row):
    """把走道上的「面上／面下」標記歸給旁邊的區塊。

    標記固定寫在某幾個走道欄位，但區塊被障礙物截短時會離得很遠，
    所以先用「緊鄰單一區塊」的標記學出「欄位 → 區域」，再套用到整欄。
    """
    marks = {(r, c): FACING[t] for (r, c), t in grid.text.items() if t in FACING}

    votes = {}
    for (r, c) in marks:
        near = sorted((block_distance(b, c), i) for i, b in enumerate(blocks_by_row.get(r, [])))
        if near and near[0][0] <= 2 and (len(near) == 1 or near[0][0] < near[1][0]):
            area = blocks_by_row[r][near[0][1]]["area"]
            votes.setdefault(c, {}).setdefault(area, 0)
            votes[c][area] += 1
    col_area = {c: max(v, key=v.get) for c, v in votes.items()}

    def side_for(block):
        best = None
        for (mr, mc), side in marks.items():
            if mr != block["r"] or col_area.get(mc) != block["area"]:
                continue
            d = block_distance(block, mc)
            if best is None or d < best[0]:
                best = (d, side)
        return best[1] if best else None

    return side_for


def parse_sheet(grid: Grid):
    seats = find_seats(grid)
    if not seats:
        return None

    # 1. 找出「座位編號列」，排除散落在說明文字裡的編號（例如表頭的「傳送門：A區A01」）
    by_row = {}
    for s in seats:
        by_row.setdefault(s["r"], []).append(s)
    id_rows = {r for r, lst in by_row.items() if len(lst) >= MIN_SEATS_PER_ID_ROW}
    by_row = {r: lst for r, lst in by_row.items() if r in id_rows}
    if not by_row:
        return None

    # 2. 每一列切成區塊（A/B/C 三區在同一列各排各的，方向也可能不同）
    blocks, blocks_by_row = [], {}
    for r in sorted(by_row):
        blocks_by_row[r] = []
        for b in split_blocks(by_row[r]):
            blk = {"r": r, "area": b[0]["area"], "c0": b[0]["c0"], "c1": b[-1]["c1"],
                   "seats": b, "side": None}
            blocks.append(blk)
            blocks_by_row[r].append(blk)

    seat_at_col, block_at_col = {}, {}
    for blk in blocks:
        for s in blk["seats"]:
            for col in range(s["c0"], s["c1"] + 1):
                seat_at_col[(blk["r"], col)] = s
                block_at_col[(blk["r"], col)] = blk

    warnings = []

    # 3. 判斷每個區塊的暱稱寫在上面還是下面
    def name_count(r, c0, c1):
        """這一列在這段欄位裡有幾格「像暱稱」的文字（排除座位編號與面上／面下）。"""
        return sum(1 for col in range(c0, c1 + 1)
                   if (t := grid.text.get((r, col))) and not SEAT_RE.match(t) and t not in FACING)

    mark_side = mark_reader(grid, blocks, blocks_by_row)
    for blk in blocks:
        r, c0, c1 = blk["r"], blk["c0"], blk["c1"]
        # 三個線索：座位編號是蛇行的，一排由小到大就是面上、由大到小就是面下；
        # 走道上的面上／面下標記；以及暱稱實際寫在哪一側。三者正常會一致。
        nums = [s["num"] for s in blk["seats"]]
        asc = sum(b > a for a, b in zip(nums, nums[1:]))
        desc = sum(b < a for a, b in zip(nums, nums[1:]))
        by_number = +1 if asc > desc else -1 if desc > asc else None
        by_mark = mark_side(blk)
        above, below = name_count(r - 1, c0, c1), name_count(r + 1, c0, c1)
        by_name = -1 if above > below else +1 if below > above else None

        votes = [v for v in (by_number, by_name, by_mark) if v is not None]
        where = f"第 {r} 列 {blk['area']} 區（{blk['seats'][0]['id']}～{blk['seats'][-1]['id']}）"
        if not votes:
            up = any((r - 1, col) in seat_at_col for col in range(c0, c1 + 1))
            dn = any((r + 1, col) in seat_at_col for col in range(c0, c1 + 1))
            blk["side"] = +1 if up and not dn else -1 if dn and not up else +1
            warnings.append(f"{where}看不出面上還是面下，方向是推測的，請人工確認")
            continue
        blk["side"] = votes[0]
        if len(set(votes)) > 1:
            told = [f"{k}={'面上' if v > 0 else '面下'}"
                    for k, v in (("編號順序", by_number), ("暱稱位置", by_name), ("原表標記", by_mark))
                    if v is not None]
            warnings.append(f"{where}的方向線索互相矛盾（{'、'.join(told)}），暫以編號順序為準，請人工確認")

    # 4. 把暱稱配給座位：一格暱稱歸給「面向它、而且離它最近」的區塊
    result = {s["id"]: {"id": s["id"], "area": s["area"], "num": s["num"],
                        "cell": s["cell"], "color": s["color"],
                        "name": "", "big": False, "partner": None}
              for s in seats if s["r"] in id_rows}

    claims = []
    for (tr, tc), t in sorted(grid.text.items()):
        if SEAT_RE.match(t) or t in FACING or ZONE_RE.match(t):
            continue
        nc0, nc1 = grid.col_span(tr, tc)
        cands, wrong_side = [], False
        for blk in blocks:
            if nc1 < blk["c0"] or nc0 > blk["c1"]:
                continue
            d = tr - blk["r"]
            if not 1 <= abs(d) <= MAX_NAME_GAP:
                continue
            if (d > 0) == (blk["side"] > 0):
                cands.append((abs(d), blk))
            else:
                wrong_side = True
        if not cands:
            if wrong_side:
                warnings.append(f"{get_column_letter(tc)}{tr}「{t}」旁邊的座位不是背對它、"
                                f"就是不在它的欄位範圍內，沒有歸給任何座位，請對照原表確認")
            continue

        best = min(d for d, _ in cands)
        tied = [b for d, b in cands if d == best]
        if len(tied) > 1:
            where = "、".join(f"第 {b['r']} 列 {b['seats'][0]['id']}" for b in tied)
            warnings.append(f"{get_column_letter(tc)}{tr}「{t}」同時貼著 {where}，請人工確認")
        claims.append((best, tied[0], nc0, nc1, t))

    # 近的先佔：原表偶爾會在暱稱下面再寫一行註記（例如「這邊應該有柱子」），
    # 那一行離座位比較遠，不該蓋掉正上／正下方的暱稱
    taken = {}
    for dist, blk, nc0, nc1, t in sorted(claims, key=lambda x: x[0]):
        covered = []
        for col in range(max(nc0, blk["c0"]), min(nc1, blk["c1"]) + 1):
            s = seat_at_col.get((blk["r"], col))
            if s and s["id"] not in covered:
                covered.append(s["id"])
        if len(covered) > MAX_BIG_SEATS:
            continue  # 橫跨一整排的多半是走道說明，不是暱稱
        for sid in covered:
            if sid in taken and taken[sid][0] == dist and taken[sid][1] != t:
                warnings.append(f"{sid} 對到兩個暱稱：{taken[sid][1]} / {t}")
        covered = [sid for sid in covered if sid not in taken]
        for sid in covered:
            result[sid]["name"] = t
            taken[sid] = (dist, t)
        if len(covered) >= 2:
            for sid in covered:
                result[sid]["big"] = True
                others = [x for x in covered if x != sid]
                result[sid]["partner"] = others[0] if len(others) == 1 else others

    # 5. 相對位置：左右鄰座、同桌對面、背後
    def seat_in_row(r, c0, c1):
        for col in range(c0, c1 + 1):
            if (r, col) in seat_at_col:
                return seat_at_col[(r, col)], block_at_col[(r, col)]
        return None, None

    area_rows = {}
    present = {}
    for blk in blocks:
        area_rows.setdefault(blk["area"], set()).add(blk["r"])
        present.setdefault(blk["area"], set()).update(s["num"] for s in blk["seats"])
    area_rows = {a: sorted(v) for a, v in area_rows.items()}

    for blk in blocks:
        blk["lo"] = min(s["num"] for s in blk["seats"])
        blk["hi"] = max(s["num"] for s in blk["seats"])

    def same_table(a, b):
        """同一張桌子的兩排，編號是接續的（中間只能跳過原表沒有的號碼）。"""
        if a["area"] != b["area"]:
            return False
        pool = present[a["area"]]
        return any(hi > lo and all(n not in pool for n in range(lo + 1, hi))
                   for lo, hi in ((a["hi"], b["lo"]), (b["hi"], a["lo"])))

    for blk in blocks:
        r, side = blk["r"], blk["side"]
        for s in blk["seats"]:
            x = result[s["id"]]
            c0, c1 = s["c0"], s["c1"]
            # 同桌對面：往暱稱那一側找第一個有座位的列，兩邊要面對面、編號也要接得上
            x["opposite"] = None
            for d in range(2, MAX_TABLE_GAP + 1):
                other, oblk = seat_in_row(r + side * d, c0, c1)
                if other:
                    if oblk["side"] == -side and same_table(blk, oblk):
                        x["opposite"] = other["id"]
                    break
            # 背後：緊貼的上／下一列，而且是背對背
            other, oblk = seat_in_row(r - side, c0, c1)
            x["behind"] = other["id"] if other and oblk["side"] == -side else None
            x["left"] = seat_at_col[(r, c0 - 1)]["id"] if block_at_col.get((r, c0 - 1)) is blk else None
            x["right"] = seat_at_col[(r, c1 + 1)]["id"] if block_at_col.get((r, c1 + 1)) is blk else None
            x["row_no"] = area_rows[blk["area"]].index(r) + 1
            x["facing"] = "面上排（暱稱寫在下方）" if side > 0 else "面下排（暱稱寫在上方）"
            x["_row"] = r

    # 6. 小地圖用：每個區塊的座位順序（照原表左到右）
    rows_out = [{"r": blk["r"], "side": blk["side"], "seats": [s["id"] for s in blk["seats"]]}
                for blk in blocks]

    seats_out = sorted(result.values(), key=lambda v: (v["area"], v["num"]))
    return {
        "generated_at": datetime.now(TZ_TAIPEI).strftime("%Y-%m-%d %H:%M"),
        "sheet": grid.ws.title,
        "total": len(seats_out),
        "occupied": sum(1 for s in seats_out if s["name"]),
        "rows": rows_out,
        "seats": seats_out,
        "zones": find_zones(grid),
        "places": find_places(grid, (min(c for _, c in seat_at_col), max(c for _, c in seat_at_col))),
        "warnings": warnings,
    }


def parse(path: str, gids: dict = None):
    wb = openpyxl.load_workbook(path, data_only=True)
    # 挑「解析出最多座位」的那張工作表；一樣多就挑登記人數多的（比較新的那張）
    best, zone_names, zone_about = None, {}, {}
    for ws in wb.worksheets:
        grid = Grid(ws)
        for z in zone_notes(grid, find_zones(grid)):   # 介紹寫在「空白領域」那張，不是地圖那張
            if z.get("about"):
                zone_about.setdefault(z["no"], z["about"])
            # 比對時忽略空白與各種破折號，只有真的換了團名才算不一致
            key = re.sub(r"[\s\-－–—~～]+", "", z["name"])
            zone_names.setdefault(z["no"], {}).setdefault(key, (z["name"], []))[1].append(ws.title)
        data = parse_sheet(grid)
        if data is None:
            continue
        key = (data["total"], data["occupied"])
        if best is None or key > best[0]:
            best = (key, data)
    if best is None:
        sys.exit("找不到任何座位編號（格式應為 A0001、B0123 這種）")
    data = best[1]
    for z in data["zones"]:
        if zone_about.get(z["no"]):
            z["about"] = zone_about[z["no"]]
    data["notes"], data["notes_sheet"] = find_notes(wb)
    data["gids"] = gids or {}
    for no, names in sorted(zone_names.items()):
        if len(names) > 1:
            told = "、".join(f"「{n}」({'/'.join(s)})" for n, s in names.values())
            data["warnings"].append(f"空白領域{no:02d} 各工作表寫的團名不一樣：{told}，請對照原表確認")
    return data


def area_summary(data):
    """每區座位數與編號斷掉的地方；原表本來就有缺號，印出來方便人工核對。"""
    lines = []
    for area in sorted({s["area"] for s in data["seats"]}):
        nums = sorted(s["num"] for s in data["seats"] if s["area"] == area)
        occupied = sum(1 for s in data["seats"] if s["area"] == area and s["name"])
        gaps = [f"{area}{a + 1:04d}" + (f"～{area}{b - 1:04d}" if b - a > 2 else "")
                for a, b in zip(nums, nums[1:]) if b - a > 1]
        lines.append(f"  {area} 區：{len(nums)} 個座位（編號到 {area}{nums[-1]:04d}）、{occupied} 個已登記"
                     + (f"；原表缺號 {'、'.join(gaps)}" if gaps else ""))
    zones = data.get("zones") or []
    if zones:
        lines.append(f"  空白領域：{len(zones)} 個營區　" +
                     "、".join(f"{z['label']} {z['name']}" for z in zones))
    return lines


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", nargs="?")
    ap.add_argument("--sheet-id")
    ap.add_argument("-o", "--out", default="seats.json")
    a = ap.parse_args()
    if not a.xlsx and not a.sheet_id:
        ap.error("請給 xlsx 檔案路徑或 --sheet-id")
    path = a.xlsx or download_xlsx(a.sheet_id)
    data = parse(path, fetch_gids(a.sheet_id) if a.sheet_id else {})
    with open(a.out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
    print(f"✔ 工作表「{data['sheet']}」：{data['total']} 個座位，{data['occupied']} 個已登記暱稱 → {a.out}")
    for line in area_summary(data):
        print(line)
    for w in data["warnings"]:
        print("⚠", w)


if __name__ == "__main__":
    main()
